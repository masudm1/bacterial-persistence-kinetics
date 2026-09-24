"""Core routines for the corrected bacterial-persistence analysis.

The raw workbooks are read without modification. Real-data fitting is done on
individual log-relative measurements, y_r(t) / y_r(0), so each biological
replicate contributes observations while sharing the same normalized initial
condition. Synthetic data use deterministic ODE trajectories plus explicitly
defined, heteroscedastic Gaussian measurement error in population space.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.optimize import least_squares


SYNTHETIC_TRUE = {"mu": 1.0, "delta": 2.5, "alpha": 0.001, "beta": 0.1}
SYNTHETIC_INITIAL = (100_000.0, 100.0)
NOISE_TIERS = {"low": 0.05, "moderate": 0.10, "high": 0.20}


@dataclass(frozen=True)
class Split:
    train_times: np.ndarray
    test_times: np.ndarray


def _valid_number(value: object) -> bool:
    return value is not None and np.isfinite(float(value)) and float(value) > 0


def extract_real_data(dryad_dir: Path) -> pd.DataFrame:
    """Return a long, replicate-level table for the two Umetani conditions.

    Columns B:D are the three ciprofloxacin replicates. Columns B:G are the
    six ampicillin replicates; missing cells remain missing and are not
    imputed. Each measurement is normalized to its own replicate's t=0 value.
    """
    specs = [
        ("ciprofloxacin", "Fig8S2.xlsx", "Fig8S2", 3),
        ("ampicillin", "Fig1S3.xlsx", "Fig1S3A", 6),
    ]
    records: list[dict[str, float | int | str]] = []
    for condition, filename, sheet, n_replicates in specs:
        workbook = load_workbook(dryad_dir / filename, read_only=True, data_only=True)
        rows = list(workbook[sheet].iter_rows(values_only=True))[4:16]
        baselines = [float(rows[0][j]) for j in range(1, n_replicates + 1)]
        for row_index, row in enumerate(rows, start=5):
            time_h = float(row[0]) / 60.0
            for replicate in range(1, n_replicates + 1):
                raw = row[replicate]
                if not _valid_number(raw):
                    continue
                relative = float(raw) / baselines[replicate - 1]
                records.append(
                    {
                        "condition": condition,
                        "source_file": filename,
                        "source_sheet": sheet,
                        "source_row": row_index,
                        "replicate": replicate,
                        "time_h": time_h,
                        "raw_cells_ml": float(raw),
                        "baseline_cells_ml": baselines[replicate - 1],
                        "relative_abundance": relative,
                        "log_relative_abundance": float(np.log(relative)),
                    }
                )
    result = pd.DataFrame.from_records(records)
    expected_amp_counts = np.array([6, 5, 6, 6, 6, 6, 3, 6, 6, 3, 4, 3])
    amp_counts = (
        result[result.condition.eq("ampicillin")]
        .groupby("time_h", sort=True)
        .size()
        .to_numpy()
    )
    if not np.array_equal(amp_counts, expected_amp_counts):
        raise ValueError(f"Ampicillin extraction failed audit: {amp_counts.tolist()}")
    return result


def extract_peyrusson_oxacillin(workbook_path: Path) -> pd.DataFrame:
    """Extract extracellular oxacillin log10-survival replicates from Fig. 1b."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = list(workbook["Fig. 1b"].iter_rows(values_only=True))
    records = []
    for source_row, row in enumerate(rows[2:7], start=3):
        for replicate, column in enumerate((10, 11, 12), start=1):
            value = row[column]
            if value is None or not np.isfinite(float(value)):
                continue
            log10_relative = float(value)
            records.append(
                {
                    "condition": "peyrusson_extracellular_oxacillin",
                    "source_file": workbook_path.name,
                    "source_sheet": "Fig. 1b",
                    "source_row": source_row,
                    "replicate": replicate,
                    "time_h": float(row[0]),
                    "log10_relative_abundance": log10_relative,
                    "log_relative_abundance": log10_relative * np.log(10.0),
                }
            )
    result = pd.DataFrame.from_records(records)
    counts = result.groupby("time_h", sort=True).size().to_numpy()
    if not np.array_equal(counts, [3, 3, 3, 3, 3]):
        raise ValueError(f"Peyrusson extraction failed audit: {counts.tolist()}")
    return result


def extract_nordholt_disinfectants(workbook_path: Path) -> pd.DataFrame:
    """Extract replicate-level Nordholt Fig. 1 source data without imputation.

    Dataset S1 reports CFU/mL and time in minutes. Each positive observation is
    normalized to its own biological replicate's time-zero count. Zero counts
    are retained and flagged in the audit table, but have no finite logarithm
    and are excluded from log-scale model fitting, matching the source paper's
    stated fitting rule.
    """
    specs = [
        ("nordholt_h2o2", "A_h2o2", 5),
        ("nordholt_bac", "D_bac", 3),
        ("nordholt_ddac", "E_ddac", 6),
    ]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    records: list[dict[str, float | int | str | bool]] = []
    for condition, sheet_name, n_replicates in specs:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        if len(rows[0]) < n_replicates + 1:
            raise ValueError(f"{sheet_name} does not contain {n_replicates} replicate columns")
        baselines = [float(rows[1][column]) for column in range(1, n_replicates + 1)]
        if any(value <= 0 or not np.isfinite(value) for value in baselines):
            raise ValueError(f"Invalid time-zero baseline in {sheet_name}")
        for source_row, row in enumerate(rows[1:], start=2):
            time_min = float(row[0])
            for replicate in range(1, n_replicates + 1):
                raw = row[replicate]
                if raw is None or not np.isfinite(float(raw)) or float(raw) < 0:
                    raise ValueError(
                        f"Invalid source value in {sheet_name}, row {source_row}, replicate {replicate}"
                    )
                raw_value = float(raw)
                relative = raw_value / baselines[replicate - 1]
                records.append(
                    {
                        "condition": condition,
                        "source_file": workbook_path.name,
                        "source_sheet": sheet_name,
                        "source_row": source_row,
                        "replicate": replicate,
                        "time_min": time_min,
                        "time_h": time_min / 60.0,
                        "raw_cells_ml": raw_value,
                        "baseline_cells_ml": baselines[replicate - 1],
                        "is_zero_count": raw_value == 0.0,
                        "relative_abundance": relative,
                        "log_relative_abundance": float(np.log(relative)) if relative > 0 else np.nan,
                    }
                )
    result = pd.DataFrame.from_records(records)
    expected = {"nordholt_h2o2": (16, 5), "nordholt_bac": (8, 3), "nordholt_ddac": (16, 6)}
    for condition, (n_times, n_replicates) in expected.items():
        subset = result[result.condition.eq(condition)]
        if subset.time_h.nunique() != n_times or subset.replicate.nunique() != n_replicates:
            raise ValueError(f"Nordholt extraction audit failed for {condition}")
        if len(subset) != n_times * n_replicates:
            raise ValueError(f"Unexpected missing cells for {condition}")
    return result


def extract_classification_anchors(dryad_dir: Path) -> dict[str, dict[str, float]]:
    """Read the two exponential-phase classification counts from Fig3.xlsx."""
    workbook = load_workbook(dryad_dir / "Fig3.xlsx", read_only=True, data_only=True)
    rows = list(workbook["Fig3"].iter_rows(values_only=True))[2:]
    targets = {
        ("LB", "Exponential", "Amp 200 µg/mL"): "ampicillin",
        ("M9", "Exponential", "CPFX 1 µg/mL"): "ciprofloxacin",
    }
    result = {}
    for row in rows:
        key = tuple(row[:3])
        if key not in targets:
            continue
        exposed = float(row[3])
        non_growing, growing, indistinguishable, total = map(float, row[4:8])
        if not np.isclose(non_growing + growing + indistinguishable, total):
            raise ValueError(f"Classification counts do not sum for {key}")
        result[targets[key]] = {
            "exposed_cells": exposed,
            "non_growing": non_growing,
            "growing": growing,
            "indistinguishable": indistinguishable,
            "total_classified_persisters": total,
            "total_fraction": total / exposed,
        }
    if set(result) != {"ampicillin", "ciprofloxacin"}:
        raise ValueError("Required Fig3 classification rows were not found")
    return result


def make_time_split(
    times: Sequence[float], seed: int, test_fraction: float = 0.20, include_zero: bool = True
) -> Split:
    """Split unique time points; all replicates at a time stay in one fold."""
    unique = np.unique(np.asarray(times, dtype=float))
    eligible = unique[unique != 0.0] if include_zero else unique
    n_test = max(1, int(round(test_fraction * len(unique))))
    n_test = min(n_test, len(eligible))
    rng = np.random.default_rng(seed)
    test = np.sort(rng.choice(eligible, size=n_test, replace=False))
    train = np.array([t for t in unique if t not in set(test)], dtype=float)
    if include_zero and 0.0 not in train:
        raise AssertionError("t=0 must remain in the training set")
    return Split(train_times=train, test_times=test)


def mask_for_times(values: Sequence[float], selected: Iterable[float]) -> np.ndarray:
    selected_array = np.asarray(tuple(selected), dtype=float)
    return np.isclose(np.asarray(values, dtype=float)[:, None], selected_array[None, :]).any(axis=1)


def rhs_real(_t: float, state: np.ndarray, D: float, alpha: float, beta: float) -> list[float]:
    N, P = state
    return [(D - alpha) * N + beta * P, alpha * N - beta * P]


def rhs_synthetic(t: float, state: np.ndarray, alpha: float, beta: float) -> list[float]:
    N, P = state
    kill = SYNTHETIC_TRUE["delta"] if 2.0 <= t <= 7.0 else 0.0
    return [
        (SYNTHETIC_TRUE["mu"] - kill - alpha) * N + beta * P,
        alpha * N - beta * P,
    ]


def solve_total(
    times: Sequence[float],
    initial: tuple[float, float],
    alpha: float,
    beta: float,
    D: float | None = None,
) -> np.ndarray:
    """Integrate at arbitrary times, handling the synthetic rate discontinuities."""
    requested = np.asarray(times, dtype=float)
    order = np.argsort(requested)
    sorted_times = requested[order]
    max_time = float(sorted_times[-1])
    rhs = (lambda t, z: rhs_synthetic(t, z, alpha, beta)) if D is None else (
        lambda t, z: rhs_real(t, z, D, alpha, beta)
    )
    boundaries = [0.0]
    if D is None:
        boundaries.extend(x for x in (2.0, 7.0) if 0.0 < x < max_time)
    boundaries.append(max_time)
    current = np.asarray(initial, dtype=float)
    collected_t = [0.0]
    collected_y = [current.copy()]
    for left, right in zip(boundaries[:-1], boundaries[1:]):
        eval_times = sorted_times[(sorted_times > left) & (sorted_times <= right)]
        if right == left:
            continue
        eval_with_right = np.unique(np.append(eval_times, right))
        solution = solve_ivp(
            rhs,
            (left, right),
            current,
            t_eval=eval_with_right,
            rtol=1e-9,
            atol=1e-11,
            method="LSODA",
        )
        if not solution.success or np.any(~np.isfinite(solution.y)):
            raise RuntimeError("ODE integration failed")
        for t_value in eval_times:
            idx = int(np.flatnonzero(np.isclose(solution.t, t_value))[0])
            collected_t.append(float(t_value))
            collected_y.append(solution.y[:, idx].copy())
        current = solution.y[:, -1]
    lookup = {round(t, 12): y for t, y in zip(collected_t, collected_y)}
    totals_sorted = np.array([lookup[round(float(t), 12)].sum() for t in sorted_times])
    totals = np.empty_like(totals_sorted)
    totals[order] = totals_sorted
    return np.maximum(totals, np.finfo(float).tiny)


def synthetic_dataset(seed: int, noise_fraction: float, times: np.ndarray | None = None) -> pd.DataFrame:
    """Deterministic truth plus N(0, [fraction * truth]^2) measurement error."""
    if times is None:
        times = np.linspace(0.0, 12.0, 100)
    clean = solve_total(
        times,
        SYNTHETIC_INITIAL,
        SYNTHETIC_TRUE["alpha"],
        SYNTHETIC_TRUE["beta"],
    )
    rng = np.random.default_rng(seed * 100 + int(round(noise_fraction * 100)))
    noisy = np.maximum(clean + rng.normal(0.0, noise_fraction * clean), 1.0)
    return pd.DataFrame(
        {
            "time_h": times,
            "latent_total": clean,
            "observed_total": noisy,
            "log_latent_total": np.log(clean),
            "log_observed_total": np.log(noisy),
        }
    )


SYNTHETIC_STARTS = np.array(
    [
        (1e-5, 1e-3),
        (1e-4, 1e-2),
        (1e-3, 1e-1),
        (1e-2, 1e0),
        (1e-1, 1e-3),
        (1e0, 1e-1),
        (1e-5, 1e0),
        (1e0, 1e-5),
        (5e-4, 5e-2),
        (5e-3, 5e-1),
        (5e-2, 5e-3),
        (5e-1, 5e-2),
    ]
)

REAL_STARTS = np.array(
    [
        (-1.0, 1e-4, 1e-2),
        (-2.0, 1e-3, 1e-1),
        (-4.0, 1e-2, 5e-1),
        (-6.0, 1e-1, 1.0),
        (-8.0, 5e-1, 2.0),
        (-10.0, 1.0, 5.0),
        (-3.0, 1e-5, 1.0),
        (-5.0, 1.0, 1e-3),
        (-12.0, 1e-3, 5.0),
        (-0.5, 1e-1, 1e-2),
        (-15.0, 1e-5, 1e-2),
        (-7.0, 1e-2, 2.0),
    ]
)


def fit_nlls_synthetic(times: np.ndarray, log_y: np.ndarray) -> dict[str, float]:
    """Fit alpha and beta to only the supplied observations."""
    best = None

    def residual(log_parameters: np.ndarray) -> np.ndarray:
        alpha, beta = np.exp(log_parameters)
        prediction = np.log(solve_total(times, SYNTHETIC_INITIAL, alpha, beta))
        return prediction - log_y

    for start in SYNTHETIC_STARTS:
        result = least_squares(
            residual,
            np.log(start),
            bounds=(np.log([1e-8, 1e-8]), np.log([10.0, 10.0])),
            max_nfev=3000,
        )
        sse = float(2.0 * result.cost)
        if best is None or sse < best[0]:
            best = (sse, result)
    assert best is not None
    alpha, beta = np.exp(best[1].x)
    return {"alpha": float(alpha), "beta": float(beta), "train_sse": best[0]}


def fit_nlls_real(
    times: np.ndarray,
    log_y: np.ndarray,
    initial_fraction: float,
    starts: np.ndarray = REAL_STARTS,
) -> dict[str, float]:
    """Fit D, alpha and beta to replicate-level log-relative observations."""
    initial = (1.0 - initial_fraction, initial_fraction)
    best = None

    def residual(parameters: np.ndarray) -> np.ndarray:
        D = -np.exp(parameters[0])
        alpha, beta = np.exp(parameters[1:])
        unique, inverse = np.unique(times, return_inverse=True)
        prediction = np.log(solve_total(unique, initial, alpha, beta, D=D))[inverse]
        return prediction - log_y

    for D0, alpha0, beta0 in starts:
        encoded = np.log([-D0, alpha0, beta0])
        result = least_squares(
            residual,
            encoded,
            bounds=(np.log([1e-3, 1e-8, 1e-8]), np.log([30.0, 20.0, 20.0])),
            max_nfev=5000,
        )
        sse = float(2.0 * result.cost)
        if best is None or sse < best[0]:
            best = (sse, result)
    assert best is not None
    D_abs, alpha, beta = np.exp(best[1].x)
    return {"D": -float(D_abs), "alpha": float(alpha), "beta": float(beta), "train_sse": best[0]}


def predict_nlls(
    times: np.ndarray,
    result: dict[str, float],
    initial: tuple[float, float],
    synthetic: bool,
) -> np.ndarray:
    return np.log(
        solve_total(
            times,
            initial,
            result["alpha"],
            result["beta"],
            D=None if synthetic else result["D"],
        )
    )


def rmse(prediction: np.ndarray, target: np.ndarray) -> float:
    return float(np.sqrt(np.mean(np.square(np.asarray(prediction) - np.asarray(target)))))
